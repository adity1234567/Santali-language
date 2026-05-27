from __future__ import annotations

import argparse
import csv
import json
import random
import re
from collections import Counter
from pathlib import Path

import openpyxl
import pandas as pd
import pdfplumber


ROOT = Path("ground truth")
OUT = Path("synthetic_ground_truth_only")
DEFAULT_TARGET = 1200
RNG = random.Random(20260527)

MOJI_CHARS = set("ÃÂÉÊÍÌÅà¦â€žœ”˜™†ºª¼½¾")
BAD_MARKERS = ["Ã", "Â", "É", "Ê", "Í", "Ì", "Å", "à¦", "â"]


def fix_mojibake(value):
    if value is None:
        return ""
    if not isinstance(value, str):
        value = str(value)
    if any(ch in value for ch in MOJI_CHARS):
        try:
            repaired = value.encode("latin1").decode("utf-8")
            before = sum(value.count(ch) for ch in BAD_MARKERS)
            after = sum(repaired.count(ch) for ch in BAD_MARKERS)
            if after <= before:
                return repaired
        except Exception:
            pass
    return value


def read_csv(name: str) -> pd.DataFrame:
    path = ROOT / name
    if not path.exists():
        workspace_path = Path(name)
        if workspace_path.exists():
            path = workspace_path
    df = pd.read_csv(path, dtype=str, keep_default_na=False, encoding="utf-8-sig")
    for col in df.columns:
        df[col] = df[col].map(fix_mojibake)
    return df


def clean_ipa(value: str) -> str:
    value = fix_mojibake(value or "")
    value = re.sub(r"^\[\s*|\s*\]$", "", value.strip())
    value = re.sub(r"\s+", " ", value)
    return value.strip(" ;,")


def clean_phrase(value: str, fallback: str = "item") -> str:
    value = fix_mojibake(value or "").strip()
    value = re.sub(r"\([^)]*\)", "", value)
    value = re.split(r"[;/]|,\s*(?:a|the|to)?\s*", value)[0].strip()
    value = re.sub(r"\s+", " ", value)
    value = value.strip(" .:-–—")
    return value or fallback


def clean_bn(value: str, fallback: str = "বস্তু") -> str:
    value = fix_mojibake(value or "").strip()
    value = re.split(r"[,;/।]", value)[0].strip()
    value = re.sub(r"\s+", " ", value)
    return value or fallback


def glossify(value: str, fallback: str = "item") -> str:
    value = clean_phrase(value, fallback).lower()
    value = re.sub(r"[^a-z0-9]+", ".", value).strip(".")
    parts = [part for part in value.split(".") if part]
    return ".".join(parts[:3]) if parts else fallback


def good_ipa(value: str, max_tokens: int = 2) -> bool:
    if not value or value in {"—", "-", "?"}:
        return False
    if len(value) > 34 or len(value.split()) > max_tokens:
        return False
    if any(ch in value for ch in "[]{};=0123456789"):
        return False
    return value.lower() not in {"none", "nan"}


def split_glosses(series) -> Counter:
    labels = Counter()
    for value in series:
        for gloss in re.split(r"\s*\|\s*|\s+", value):
            if not gloss:
                continue
            labels[gloss] += 1
            for part in re.split(r"[-.=+/]", gloss):
                if re.search(r"[A-Z0-9]", part):
                    labels[part] += 1
    return labels


def article(word: str) -> str:
    return "an" if word[:1].lower() in "aeiou" else "a"


def cap(value: str) -> str:
    return value[:1].upper() + value[1:] if value else value


def en(item) -> str:
    return item.get("en") or "item"


def bn(item) -> str:
    return item.get("bn") or "বস্তু"


BN_LOC_OVERRIDES = {
    "ɔɽaʔ": "বাড়িতে",
    "iskul": "স্কুলে",
    "at̪o": "গ্রামে",
    "sɔhɔɽ": "শহরে",
    "gaɖa": "নদীতে",
    "ɟɔla": "হ্রদে",
    "kuɲ": "কুয়ায়",
}


VERB_FORMS = {
    "ɲɛl": {"present3": "sees", "past": "saw", "part": "seen", "ing": "seeing"},
    "ɟɔm": {"present3": "eats", "past": "ate", "part": "eaten", "ing": "eating"},
    "ɲu": {"present3": "drinks", "past": "drank", "part": "drunk", "ing": "drinking"},
    "ɾɔɽ": {"present3": "speaks", "past": "spoke", "part": "spoken", "ing": "speaking"},
    "ɔlɔʔ": {"present3": "writes", "past": "wrote", "part": "written", "ing": "writing"},
    "paɽha͡o̯": {"present3": "studies", "past": "studied", "part": "studied", "ing": "studying"},
    "kiɾiɲ": {"present3": "buys", "past": "bought", "part": "bought", "ing": "buying"},
    "ɲam": {"present3": "gets", "past": "got", "part": "gotten", "ing": "getting"},
    "ɘkʰɾiɲ": {"present3": "sells", "past": "sold", "part": "sold", "ing": "selling"},
    "gɔɽɔ": {"present3": "helps", "past": "helped", "part": "helped", "ing": "helping"},
    "calaʔ": {"present3": "goes", "past": "went", "part": "gone", "ing": "going"},
    "hiɟuʔ": {"present3": "comes", "past": "came", "part": "come", "ing": "coming"},
    "d̪uɽup̚ʔ": {"present3": "sits", "past": "sat", "part": "sat", "ing": "sitting"},
    "kɘmi": {"present3": "works", "past": "worked", "part": "worked", "ing": "working"},
    "t̪ɘŋgi": {"present3": "waits", "past": "waited", "part": "waited", "ing": "waiting"},
    "gat̪e": {"present3": "plays", "past": "played", "part": "played", "ing": "playing"},
    "ɾaːg": {"present3": "sounds", "past": "sounded", "part": "sounded", "ing": "sounding"},
    "gɔc̚ʔ": {"present3": "dies", "past": "died", "part": "died", "ing": "dying"},
}

BN_VERB_FORMS = {
    "ɲɛl": {"future": "দেখবে", "past": "দেখেছিল", "perfect": "দেখেছে", "prog": "দেখছে", "cvb": "দেখে", "imp": "দেখো", "habit": "দেখে"},
    "ɟɔm": {"future": "খাবে", "past": "খেয়েছিল", "perfect": "খেয়েছে", "prog": "খাচ্ছে", "cvb": "খেয়ে", "imp": "খাও", "habit": "খায়"},
    "ɲu": {"future": "পান করবে", "past": "পান করেছিল", "perfect": "পান করেছে", "prog": "পান করছে", "cvb": "পান করে", "imp": "পান করো", "habit": "পান করে"},
    "ɾɔɽ": {"future": "বলবে", "past": "বলেছিল", "perfect": "বলেছে", "prog": "বলছে", "cvb": "বলে", "imp": "বলো", "habit": "বলে"},
    "ɔlɔʔ": {"future": "লিখবে", "past": "লিখেছিল", "perfect": "লিখেছে", "prog": "লিখছে", "cvb": "লিখে", "imp": "লিখো", "habit": "লিখে"},
    "paɽha͡o̯": {"future": "পড়বে", "past": "পড়েছিল", "perfect": "পড়েছে", "prog": "পড়ছে", "cvb": "পড়ে", "imp": "পড়ো", "habit": "পড়ে"},
    "kiɾiɲ": {"future": "কিনবে", "past": "কিনেছিল", "perfect": "কিনেছে", "prog": "কিনছে", "cvb": "কিনে", "imp": "কিনো", "habit": "কেনে"},
    "ɲam": {"future": "পাবে", "past": "পেয়েছিল", "perfect": "পেয়েছে", "prog": "পাচ্ছে", "cvb": "পেয়ে", "imp": "নাও", "habit": "পায়"},
    "ɘkʰɾiɲ": {"future": "বিক্রি করবে", "past": "বিক্রি করেছিল", "perfect": "বিক্রি করেছে", "prog": "বিক্রি করছে", "cvb": "বিক্রি করে", "imp": "বিক্রি করো", "habit": "বিক্রি করে"},
    "gɔɽɔ": {"future": "সাহায্য করবে", "past": "সাহায্য করেছিল", "perfect": "সাহায্য করেছে", "prog": "সাহায্য করছে", "cvb": "সাহায্য করে", "imp": "সাহায্য করো", "habit": "সাহায্য করে"},
    "calaʔ": {"future": "যাবে", "past": "গিয়েছিল", "perfect": "গিয়েছে", "prog": "যাচ্ছে", "cvb": "গিয়ে", "imp": "যাও", "habit": "যায়"},
    "hiɟuʔ": {"future": "আসবে", "past": "এসেছিল", "perfect": "এসেছে", "prog": "আসছে", "cvb": "এসে", "imp": "আসো", "habit": "আসে"},
    "d̪uɽup̚ʔ": {"future": "বসবে", "past": "বসেছিল", "perfect": "বসেছে", "prog": "বসছে", "cvb": "বসে", "imp": "বসো", "habit": "বসে"},
    "kɘmi": {"future": "কাজ করবে", "past": "কাজ করেছিল", "perfect": "কাজ করেছে", "prog": "কাজ করছে", "cvb": "কাজ করে", "imp": "কাজ করো", "habit": "কাজ করে"},
    "t̪ɘŋgi": {"future": "অপেক্ষা করবে", "past": "অপেক্ষা করেছিল", "perfect": "অপেক্ষা করেছে", "prog": "অপেক্ষা করছে", "cvb": "অপেক্ষা করে", "imp": "অপেক্ষা করো", "habit": "অপেক্ষা করে"},
    "gat̪e": {"future": "খেলবে", "past": "খেলেছিল", "perfect": "খেলেছে", "prog": "খেলছে", "cvb": "খেলে", "imp": "খেলো", "habit": "খেলে"},
    "ɾaːg": {"future": "শব্দ করবে", "past": "শব্দ করেছিল", "perfect": "শব্দ করেছে", "prog": "শব্দ করছে", "cvb": "শব্দ করে", "imp": "শব্দ করো", "habit": "শব্দ করে"},
    "gɔc̚ʔ": {"future": "মারা যাবে", "past": "মারা গিয়েছিল", "perfect": "মারা গিয়েছে", "prog": "মারা যাচ্ছে", "cvb": "মরে", "imp": "মরো", "habit": "মরে"},
}


def bn_loc(item) -> str:
    return BN_LOC_OVERRIDES.get(item.get("ipa"), f"{bn(item)}ে")


def subject_text(pronoun, lower: bool = False) -> str:
    value = pronoun["en"]
    if lower and value != "I":
        return value.lower()
    return value


def subject_start(pronoun) -> str:
    return cap(subject_text(pronoun))


def do_aux(pronoun) -> str:
    return "does" if pronoun["gl"] == "3SG.ANIM" else "do"


def be_aux(pronoun) -> str:
    if pronoun["gl"] == "1SG":
        return "am"
    if pronoun["gl"] in {"1PL.EXCL", "2PL", "3PL.ANIM"}:
        return "are"
    return "is"


def verb_form(verb, form: str) -> str:
    forms = VERB_FORMS.get(verb["ipa"], {})
    if form in forms:
        return forms[form]
    base = verb["en"]
    if form == "ing":
        return base + "ing"
    if form in {"past", "part"}:
        return base + "ed"
    if form == "present3":
        return base + "s"
    return base


def bn_verb(verb, form: str) -> str:
    return BN_VERB_FORMS.get(verb["ipa"], {}).get(form, verb.get("bn", "করা"))


def bn_person_verb(verb, form: str, pronoun) -> str:
    value = bn_verb(verb, form)
    gl = pronoun["gl"]
    if form == "future" and gl in {"1SG", "1PL.EXCL"} and value.endswith("বে"):
        return value[:-2] + "ব"
    if form in {"prog", "perfect"}:
        if gl in {"1SG", "1PL.EXCL"} and value.endswith("ছে"):
            return value[:-2] + "ছি"
        if gl in {"2SG", "2PL"} and value.endswith("ছে"):
            return value[:-2] + "ছ"
    if form == "past":
        if gl in {"1SG", "1PL.EXCL"}:
            return value + "াম"
        if gl in {"2SG", "2PL"}:
            return value + "ে"
    return value


def place_phrase(verb, place) -> str:
    prep = "to" if verb["ipa"] in {"calaʔ", "hiɟuʔ"} else "in"
    return f"{prep} the {en(place)}"


def pick(seq):
    return RNG.choice(seq)


def surface_word(segment: str) -> str:
    return segment.replace("-", "")


def surface(segments) -> str:
    text = " ".join(surface_word(seg) for seg in segments)
    text = re.sub(r"\s+", " ", text).strip()
    return text if text.endswith((".", "?", "!")) else text + "."


def normalized(segments) -> str:
    return " ".join(segments).strip() + "."


def read_ground_truth():
    """Read all files from ground truth only; do not inspect generated."""
    layout_a = read_csv("output_layoutA (1).csv")
    layout_b = read_csv("output_layoutB (1).csv")
    parallel = read_csv("Santali_dataset_parsed_bom.csv")
    parallel_raw = read_csv("Santali_dataset_parsed.csv")

    text_files = {
        path.name: fix_mojibake(path.read_text(encoding="utf-8", errors="replace"))
        for path in sorted(ROOT.glob("*.txt"))
    }

    pdf_texts = {}
    pdf_meta = {}
    for path in sorted(ROOT.glob("*.pdf")):
        with pdfplumber.open(path) as pdf:
            pdf_meta[path.name] = dict(pdf.metadata or {}) | {"pages": len(pdf.pages)}
            chunks = []
            for page in pdf.pages:
                try:
                    chunks.append(page.extract_text() or "")
                except Exception:
                    chunks.append("")
            pdf_texts[path.name] = fix_mojibake("\n".join(chunks))

    wb = openpyxl.load_workbook(ROOT / "MTZ_Santali_Dictionary.xlsx", read_only=True, data_only=True)
    ws = wb["Dictionary"]
    dictionary_rows = []
    header = None
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        values = [fix_mojibake(value) for value in row]
        if i == 0:
            header = values
        else:
            dictionary_rows.append(dict(zip(header, values)))

    semantic_domains = []
    sem_ws = wb["Semantic_Domains"]
    for i, row in enumerate(sem_ws.iter_rows(values_only=True)):
        if i == 0:
            continue
        values = [fix_mojibake(value) for value in row]
        if len(values) >= 2 and values[1]:
            semantic_domains.append(values[1])

    return layout_a, layout_b, parallel, parallel_raw, text_files, pdf_texts, pdf_meta, dictionary_rows, header, semantic_domains


CORE_NOUNS = [
    {"ipa": "hɔɽ", "en": "person", "bn": "মানুষ", "gloss": "person", "domain": "Person", "class": "human"},
    {"ipa": "gid̪ɾɘ", "en": "child", "bn": "শিশু", "gloss": "child", "domain": "Person", "class": "human"},
    {"ipa": "gɔgɔ", "en": "mother", "bn": "মা", "gloss": "mother", "domain": "Kinship", "class": "human"},
    {"ipa": "baba", "en": "father", "bn": "বাবা", "gloss": "father", "domain": "Kinship", "class": "human"},
    {"ipa": "bɔ͡i̯ha", "en": "brother", "bn": "ভাই", "gloss": "brother", "domain": "Kinship", "class": "human"},
    {"ipa": "ɖak̚t̪aɾ", "en": "doctor", "bn": "ডাক্তার", "gloss": "doctor", "domain": "Health", "class": "human"},
    {"ipa": "ɟɘliɘ", "en": "fisherman", "bn": "জেলে", "gloss": "fisherman", "domain": "Person", "class": "human"},
    {"ipa": "ɔɽaʔ", "en": "house", "bn": "বাড়ি", "gloss": "house", "domain": "Place", "class": "place"},
    {"ipa": "iskul", "en": "school", "bn": "স্কুল", "gloss": "school", "domain": "Education", "class": "place"},
    {"ipa": "at̪o", "en": "village", "bn": "গ্রাম", "gloss": "village", "domain": "Place", "class": "place"},
    {"ipa": "sɔhɔɽ", "en": "town", "bn": "শহর", "gloss": "town", "domain": "Place", "class": "place"},
    {"ipa": "gaɖa", "en": "river", "bn": "নদী", "gloss": "river", "domain": "Nature", "class": "place"},
    {"ipa": "ɟɔla", "en": "lake", "bn": "হ্রদ", "gloss": "lake", "domain": "Nature", "class": "place"},
    {"ipa": "kuɲ", "en": "well", "bn": "কুয়া", "gloss": "well", "domain": "Place", "class": "place"},
    {"ipa": "pɘɾsi", "en": "language", "bn": "ভাষা", "gloss": "language", "domain": "Speech", "class": "abstract"},
    {"ipa": "kat̪ʰa", "en": "word", "bn": "কথা", "gloss": "word", "domain": "Speech", "class": "abstract"},
    {"ipa": "ɟiŋgi", "en": "life", "bn": "জীবন", "gloss": "life", "domain": "Abstract", "class": "abstract"},
    {"ipa": "d̪aʔ", "en": "water", "bn": "পানি", "gloss": "water", "domain": "Nature", "class": "liquid"},
    {"ipa": "t̪o̯͡a", "en": "milk", "bn": "দুধ", "gloss": "milk", "domain": "Food", "class": "liquid"},
    {"ipa": "ɟɔmaʔ", "en": "food", "bn": "খাবার", "gloss": "food", "domain": "Food", "class": "food"},
    {"ipa": "ɲɛlɛɾasa", "en": "honey", "bn": "মধু", "gloss": "honey", "domain": "Food", "class": "food"},
    {"ipa": "haku", "en": "fish", "bn": "মাছ", "gloss": "fish", "domain": "Food", "class": "food"},
    {"ipa": "basaŋ", "en": "cooked food", "bn": "রান্না", "gloss": "cooked.food", "domain": "Food", "class": "food"},
    {"ipa": "ɾan", "en": "medicine", "bn": "ঔষধ", "gloss": "medicine", "domain": "Health", "class": "object"},
    {"ipa": "ɾu̯͡ɘ", "en": "illness", "bn": "রোগ", "gloss": "illness", "domain": "Health", "class": "abstract"},
    {"ipa": "mɔbaĭl pʰon", "en": "mobile phone", "bn": "মোবাইল ফোন", "gloss": "mobile.phone", "domain": "Object", "class": "object"},
    {"ipa": "pʰɔʈɔ", "en": "photo", "bn": "ছবি", "gloss": "photo", "domain": "Object", "class": "object"},
    {"ipa": "gʰuɽi", "en": "watch", "bn": "ঘড়ি", "gloss": "watch", "domain": "Object", "class": "object"},
    {"ipa": "ʈaka", "en": "money", "bn": "টাকা", "gloss": "money", "domain": "Finance", "class": "object"},
    {"ipa": "pɔɾɔb", "en": "festival", "bn": "উৎসব", "gloss": "festival", "domain": "Social", "class": "abstract"},
    {"ipa": "seɾeɲ", "en": "song", "bn": "গান", "gloss": "song", "domain": "Sound", "class": "abstract"},
    {"ipa": "akʰɾa", "en": "gathering", "bn": "আখড়া", "gloss": "gathering", "domain": "Social", "class": "abstract"},
    {"ipa": "kɘmi", "en": "work", "bn": "কাজ", "gloss": "work", "domain": "Action", "class": "abstract"},
    {"ipa": "sikʰnɘt̪", "en": "lesson", "bn": "পাঠ", "gloss": "lesson", "domain": "Education", "class": "abstract"},
    {"ipa": "ɔkt̪ɛ", "en": "time", "bn": "সময়", "gloss": "time", "domain": "Time", "class": "abstract"},
    {"ipa": "muc̚ʔ", "en": "ant", "bn": "পিঁপড়া", "gloss": "ant", "domain": "Animal", "class": "animal"},
    {"ipa": "t̪ɘɾup̚ʔ", "en": "tiger", "bn": "বাঘ", "gloss": "tiger", "domain": "Animal", "class": "animal"},
    {"ipa": "suk̚ɾi", "en": "pig", "bn": "শুকর", "gloss": "pig", "domain": "Animal", "class": "animal"},
    {"ipa": "pusi", "en": "cat", "bn": "বিড়াল", "gloss": "cat", "domain": "Animal", "class": "animal"},
    {"ipa": "kaʈa", "en": "foot", "bn": "পা", "gloss": "foot", "domain": "Body", "class": "body"},
    {"ipa": "t̪i", "en": "hand", "bn": "হাত", "gloss": "hand", "domain": "Body", "class": "body"},
    {"ipa": "ɟaŋga", "en": "leg", "bn": "পা", "gloss": "leg", "domain": "Body", "class": "body"},
]

CORE_ADJS = [
    {"ipa": "mɔ̃ɟ", "en": "good", "bn": "ভালো", "gloss": "good", "domain": "Quality"},
    {"ipa": "bʰage", "en": "good", "bn": "ভালো", "gloss": "good", "domain": "Quality"},
    {"ipa": "sibil", "en": "tasty", "bn": "সুস্বাদু", "gloss": "tasty", "domain": "Quality"},
    {"ipa": "ɟɘɾuɽ", "en": "important", "bn": "গুরুত্বপূর্ণ", "gloss": "important", "domain": "Quality"},
    {"ipa": "lɘʈu", "en": "small", "bn": "ছোট", "gloss": "small", "domain": "Quality"},
    {"ipa": "ɾɘskɘ", "en": "happy", "bn": "আনন্দের", "gloss": "happy", "domain": "Emotion"},
    {"ipa": "bɔnd̪o", "en": "closed", "bn": "বন্ধ", "gloss": "closed", "domain": "State"},
    {"ipa": "sapʰa", "en": "clean", "bn": "পরিষ্কার", "gloss": "clean", "domain": "Quality"},
    {"ipa": "nɔwa", "en": "new", "bn": "নতুন", "gloss": "new", "domain": "Quality"},
    {"ipa": "pɘhil", "en": "first", "bn": "প্রথম", "gloss": "first", "domain": "Quantity"},
    {"ipa": "a͡i̯ma", "en": "many", "bn": "অনেক", "gloss": "many", "domain": "Quantity"},
]

CORE_VERBS = [
    {"ipa": "ɲɛl", "en": "see", "bn": "দেখা", "gloss": "see", "type": "trans"},
    {"ipa": "ɟɔm", "en": "eat", "bn": "খাওয়া", "gloss": "eat", "type": "trans", "objects": ["food"]},
    {"ipa": "ɲu", "en": "drink", "bn": "পান করা", "gloss": "drink", "type": "trans", "objects": ["liquid"]},
    {"ipa": "ɾɔɽ", "en": "speak", "bn": "কথা বলা", "gloss": "speak", "type": "trans", "objects": ["abstract", "speech"]},
    {"ipa": "ɔlɔʔ", "en": "write", "bn": "লেখা", "gloss": "write", "type": "trans", "objects": ["abstract", "object"]},
    {"ipa": "paɽha͡o̯", "en": "study", "bn": "পড়া", "gloss": "study", "type": "trans", "objects": ["education", "abstract"]},
    {"ipa": "kiɾiɲ", "en": "buy", "bn": "কেনা", "gloss": "buy", "type": "trans", "objects": ["object", "food"]},
    {"ipa": "ɲam", "en": "get", "bn": "পাওয়া", "gloss": "get", "type": "trans"},
    {"ipa": "ɘkʰɾiɲ", "en": "sell", "bn": "বিক্রি করা", "gloss": "sell", "type": "trans", "objects": ["food", "object"]},
    {"ipa": "gɔɽɔ", "en": "help", "bn": "সাহায্য করা", "gloss": "help", "type": "trans", "objects": ["human", "animal"]},
    {"ipa": "calaʔ", "en": "go", "bn": "যাওয়া", "gloss": "go", "type": "intrans"},
    {"ipa": "hiɟuʔ", "en": "come", "bn": "আসা", "gloss": "come", "type": "intrans"},
    {"ipa": "d̪uɽup̚ʔ", "en": "sit", "bn": "বসা", "gloss": "sit", "type": "intrans"},
    {"ipa": "kɘmi", "en": "work", "bn": "কাজ করা", "gloss": "work", "type": "intrans"},
    {"ipa": "t̪ɘŋgi", "en": "wait", "bn": "অপেক্ষা করা", "gloss": "wait", "type": "intrans"},
    {"ipa": "gat̪e", "en": "play", "bn": "খেলা", "gloss": "play", "type": "intrans"},
    {"ipa": "ɾaːg", "en": "sound", "bn": "শব্দ করা", "gloss": "sound", "type": "intrans"},
    {"ipa": "gɔc̚ʔ", "en": "die", "bn": "মারা যাওয়া", "gloss": "die", "type": "intrans"},
]

PRONOUNS = [
    {"ipa": "iɲ", "gl": "1SG", "en": "I", "bn": "আমি", "agr": "iɲ", "neg": "bɘ-ɲ", "poss": "iɲ-aʔ", "poss_en": "my", "poss_bn": "আমার", "poss_gl": "1SG-GEN"},
    {"ipa": "am", "gl": "2SG", "en": "you", "bn": "তুমি", "agr": "ɛm", "neg": "ba-m", "poss": "am-aʔ", "poss_en": "your", "poss_bn": "তোমার", "poss_gl": "2SG-GEN"},
    {"ipa": "uni", "gl": "3SG.ANIM", "en": "he/she", "bn": "সে", "agr": "i̯", "neg": "ba-i̯", "poss": "uni-aʔ", "poss_en": "his/her", "poss_bn": "তার", "poss_gl": "3SG-GEN"},
    {"ipa": "alɛ", "gl": "1PL.EXCL", "en": "we", "bn": "আমরা", "agr": "lɛ", "neg": "ba-lɛ", "poss": "alɛ-aʔ", "poss_en": "our", "poss_bn": "আমাদের", "poss_gl": "1PL.EXCL-GEN"},
    {"ipa": "apɛ", "gl": "2PL", "en": "you all", "bn": "তোমরা", "agr": "pɛ", "neg": "ba-pɛ", "poss": "apɛ-aʔ", "poss_en": "your", "poss_bn": "তোমাদের", "poss_gl": "2PL-GEN"},
    {"ipa": "oŋko", "gl": "3PL.ANIM", "en": "they", "bn": "তারা", "agr": "ko", "neg": "ba-ko", "poss": "oŋko-aʔ", "poss_en": "their", "poss_bn": "তাদের", "poss_gl": "3PL-GEN"},
]


def build_lexicon(dictionary_rows):
    seen = {item["ipa"] for item in CORE_NOUNS + CORE_ADJS + CORE_VERBS}
    dict_nouns, dict_adjs, dict_verbs = [], [], []

    for entry in dictionary_rows:
        ipa = clean_ipa(entry.get("IPA") or entry.get("Santali (Headword)"))
        pos = (entry.get("POS") or "").lower()
        english = clean_phrase(entry.get("English"), "item").lower()
        bangla = clean_bn(entry.get("Bangla"), "বস্তু")
        domain = entry.get("Semantic Domain") or "General"

        if not good_ipa(ipa, max_tokens=2) or ipa in seen:
            continue

        item = {"ipa": ipa, "en": english, "bn": bangla, "gloss": glossify(english), "domain": domain}
        if pos.startswith("n") and len(dict_nouns) < 420:
            cls = "human" if domain in {"Person", "Kinship"} else "place" if domain == "Place" else "food" if domain == "Food" else "animal" if domain == "Animal" else "object"
            item["class"] = cls
            dict_nouns.append(item)
            seen.add(ipa)
        elif pos.startswith("adj") and len(dict_adjs) < 220:
            dict_adjs.append(item)
            seen.add(ipa)
        elif pos.startswith("v") and len(dict_verbs) < 120 and len(ipa.split()) == 1:
            item["type"] = "trans"
            dict_verbs.append(item)
            seen.add(ipa)

    all_nouns = CORE_NOUNS + dict_nouns
    all_adjs = CORE_ADJS + dict_adjs
    all_verbs = CORE_VERBS + dict_verbs
    return all_nouns, all_adjs, all_verbs, dict_nouns, dict_adjs, dict_verbs


def make_generator(all_nouns, all_adjs, all_verbs, dict_nouns, dict_verbs):
    places_core = [n for n in CORE_NOUNS if n["class"] == "place"]
    food_core = [n for n in CORE_NOUNS if n["class"] in {"food", "liquid"}]
    human_core = [n for n in CORE_NOUNS if n["class"] == "human"]
    object_core = [n for n in CORE_NOUNS if n["class"] in {"object", "abstract", "food", "liquid", "animal", "body"}]
    places = places_core + [n for n in dict_nouns if n.get("class") == "place"]
    objects = object_core + [n for n in dict_nouns if n.get("class") in {"object", "animal", "food"}]
    humans = human_core + [n for n in dict_nouns if n.get("class") == "human"]
    trans_verbs = [v for v in CORE_VERBS if v.get("type") == "trans"]
    intrans_verbs = [v for v in CORE_VERBS if v.get("type") == "intrans"]
    adj_by_ipa = {adj["ipa"]: adj for adj in CORE_ADJS}
    noun_by_ipa = {noun["ipa"]: noun for noun in all_nouns}
    special_objects = {
        "ɲu": ["d̪aʔ", "t̪o̯͡a"],
        "ɟɔm": ["ɟɔmaʔ", "haku", "basaŋ", "ɲɛlɛɾasa"],
        "ɾɔɽ": ["pɘɾsi", "kat̪ʰa"],
        "ɔlɔʔ": ["kat̪ʰa", "pɘɾsi"],
        "paɽha͡o̯": ["pɘɾsi", "kat̪ʰa", "sikʰnɘt̪"],
        "kiɾiɲ": ["gʰuɽi", "mɔbaĭl pʰon", "ɾan", "ɟɔmaʔ", "haku"],
        "ɲam": ["ʈaka", "ɾan", "pʰɔʈɔ", "ɟɔmaʔ"],
        "ɘkʰɾiɲ": ["haku", "ɟɔmaʔ", "gʰuɽi", "mɔbaĭl pʰon"],
        "gɔɽɔ": ["hɔɽ", "gid̪ɾɘ", "gɔgɔ", "baba", "ɖak̚t̪aɾ"],
    }

    def compatible_adj(noun):
        cls = noun.get("class")
        domain = noun.get("domain")
        if cls in {"food", "liquid"}:
            choices = ["sibil", "mɔ̃ɟ", "sapʰa"]
        elif cls in {"human", "animal"}:
            choices = ["mɔ̃ɟ", "lɘʈu", "ɾɘskɘ"]
        elif cls == "place":
            choices = ["mɔ̃ɟ", "sapʰa", "nɔwa", "ɟɘɾuɽ"]
        elif domain in {"Education", "Speech", "Social", "Health", "Abstract"}:
            choices = ["ɟɘɾuɽ", "mɔ̃ɟ", "nɔwa"]
        else:
            choices = ["mɔ̃ɟ", "ɟɘɾuɽ", "nɔwa", "sapʰa"]
        return pick([adj_by_ipa[ipa] for ipa in choices if ipa in adj_by_ipa])

    def pick_object_for(verb):
        if verb["ipa"] in special_objects:
            return pick([noun_by_ipa[ipa] for ipa in special_objects[verb["ipa"]] if ipa in noun_by_ipa])
        allowed = verb.get("objects")
        if not allowed:
            return pick(objects)
        pools = []
        for cls in allowed:
            if cls == "speech":
                pools.extend([n for n in all_nouns if n["domain"] == "Speech"])
            elif cls == "education":
                pools.extend([n for n in all_nouns if n["domain"] == "Education"])
            elif cls == "abstract":
                pools.extend([n for n in all_nouns if n.get("class") == "abstract"])
            else:
                pools.extend([n for n in all_nouns if n.get("class") == cls])
        return pick(pools or objects)

    def note_for(pattern):
        if RNG.random() >= 0.22:
            return ""
        notes = {
            "topic_adjective": "Synthetic TOP + adjectival-FIN clause following layout A/B labels.",
            "possessive_adjective": "Synthetic GEN + TOP clause using pronoun markers from ground truth.",
            "locative_exist": "Synthetic LOC + exist-FIN clause matching mena-a patterns.",
            "transitive_fin": "Synthetic transitive finite clause with TOP and FIN labels.",
            "negative_transitive": "Synthetic NEG-pronominal marker pattern from the see paradigm.",
            "progressive": "Synthetic PROG-FIN agreement pattern based on see-paradigm rows.",
            "past_transitive": "Synthetic PST-FIN transitive pattern using dictionary/core lexicon.",
            "perfect_transitive": "Synthetic PRF-FIN clause modeled on akad-a rows.",
            "coordination": "Synthetic coordination with aɾ 'and'.",
            "contrast": "Synthetic contrast with mɛnkʰan 'but' and NEG marking.",
            "imperative": "Synthetic polite imperative with d̪a.ja kat̪ɛ and -mɛ.",
            "question": "Synthetic question row using cɛt̪ʔ/ja patterns.",
            "sequence": "Synthetic sequential CVB clause with -kat̪ɛ.",
            "classifier_exist": "Synthetic classifier + LOC/exist clause.",
        }
        return notes.get(pattern, "Synthetic ground-truth-pattern row.")

    def row(pattern, segments, glosses, english, bangla, file_hint):
        if len(segments) != len(glosses):
            raise ValueError((pattern, segments, glosses))
        return {
            "pattern": pattern,
            "file": file_hint,
            "segments_list": segments,
            "gloss_list": glosses,
            "ipa_original": surface(segments),
            "ipa_normalized": normalized(segments),
            "segments": " | ".join(segments),
            "glosses": " | ".join(glosses),
            "english": english,
            "bangla_plain": bangla,
            "notes": note_for(pattern),
        }

    def t_topic_adjective():
        noun = pick(all_nouns if RNG.random() > 0.25 else CORE_NOUNS)
        adj = compatible_adj(noun)
        return row(
            "topic_adjective",
            [noun["ipa"], "d̪ɔ", f"{adj['ipa']}-a"],
            [noun["gloss"], "TOP", f"{adj['gloss']}-FIN"],
            f"The {en(noun)} is {en(adj)}.",
            f"{bn(noun)}টি {bn(adj)}।",
            "synthetic_texts_glossed.eaf",
        )

    def t_possessive_adjective():
        pr = pick(PRONOUNS)
        noun = pick(CORE_NOUNS if RNG.random() < 0.6 else all_nouns)
        adj = compatible_adj(noun)
        return row(
            "possessive_adjective",
            [pr["poss"], noun["ipa"], "d̪ɔ", f"{adj['ipa']}-a"],
            [pr["poss_gl"], noun["gloss"], "TOP", f"{adj['gloss']}-FIN"],
            f"{cap(pr['poss_en'])} {en(noun)} is {en(adj)}.",
            f"{pr['poss_bn']} {bn(noun)} {bn(adj)}।",
            "synthetic_paradigm_glossed.eaf",
        )

    def t_locative_exist():
        place = pick(places or places_core)
        noun = pick(CORE_NOUNS if RNG.random() < 0.35 else all_nouns)
        if noun == place:
            noun = pick(objects)
        return row(
            "locative_exist",
            [f"{place['ipa']}-ɾɛ", noun["ipa"], "mena-a"],
            [f"{place['gloss']}-LOC", noun["gloss"], "exist-FIN"],
            f"There is {article(en(noun))} {en(noun)} in the {en(place)}.",
            f"{bn_loc(place)} একটি {bn(noun)} আছে।",
            "synthetic_texts_glossed.eaf",
        )

    def t_transitive_fin():
        pr = pick(PRONOUNS)
        verb = pick(trans_verbs)
        obj = pick_object_for(verb)
        return row(
            "transitive_fin",
            [pr["ipa"], "d̪ɔ", obj["ipa"], f"{verb['ipa']}-a"],
            [pr["gl"], "TOP", obj["gloss"], f"{verb['gloss']}-FIN"],
            f"{subject_start(pr)} will {verb['en']} the {en(obj)}.",
            f"{pr['bn']} {bn(obj)} {bn_person_verb(verb, 'future', pr)}।",
            "synthetic_bangla_glossed.eaf",
        )

    def t_negative_transitive():
        pr = pick(PRONOUNS)
        verb = pick([v for v in CORE_VERBS if v["type"] == "trans"])
        obj = pick_object_for(verb)
        return row(
            "negative_transitive",
            [pr["ipa"], "d̪ɔ", obj["ipa"], pr["neg"], f"{verb['ipa']}-a"],
            [pr["gl"], "TOP", obj["gloss"], f"NEG-{pr['gl']}", f"{verb['gloss']}-FIN"],
            f"{subject_start(pr)} {do_aux(pr)} not {verb['en']} the {en(obj)}.",
            f"{pr['bn']} {bn(obj)} {bn_person_verb(verb, 'future', pr)} না।",
            "synthetic_see_paradigm.eaf",
        )

    def t_progressive():
        pr = pick(PRONOUNS)
        verb = pick(intrans_verbs if RNG.random() < 0.55 else CORE_VERBS)
        return row(
            "progressive",
            [f"{pr['ipa']}-{pr['agr']}", f"{verb['ipa']}-ɛ-d̪a"],
            [f"{pr['gl']}-{pr['gl']}", f"{verb['gloss']}-PROG-FIN"],
            f"{subject_start(pr)} {be_aux(pr)} {verb_form(verb, 'ing')}.",
            f"{pr['bn']} {bn_person_verb(verb, 'prog', pr)}।",
            "synthetic_see_paradigm.eaf",
        )

    def t_past_transitive():
        pr = pick(PRONOUNS)
        verb = pick(trans_verbs)
        obj = pick_object_for(verb)
        return row(
            "past_transitive",
            [f"{pr['ipa']}-{pr['agr']}", obj["ipa"], f"{verb['ipa']}-ked-a"],
            [f"{pr['gl']}-{pr['gl']}", obj["gloss"], f"{verb['gloss']}-PST-FIN"],
            f"{subject_start(pr)} {verb_form(verb, 'past')} the {en(obj)}.",
            f"{pr['bn']} {bn(obj)} {bn_person_verb(verb, 'past', pr)}।",
            "synthetic_paradigm_glossed.eaf",
        )

    def t_perfect_transitive():
        pr = pick(PRONOUNS)
        verb = pick([v for v in CORE_VERBS if v["type"] == "trans"])
        obj = pick_object_for(verb)
        return row(
            "perfect_transitive",
            [pr["ipa"], "d̪ɔ", obj["ipa"], f"{verb['ipa']}-akad-a"],
            [pr["gl"], "TOP", obj["gloss"], f"{verb['gloss']}-PRF-FIN"],
            f"{subject_start(pr)} {'has' if pr['gl'] == '3SG.ANIM' else 'have'} {verb_form(verb, 'part')} the {en(obj)}.",
            f"{pr['bn']} {bn(obj)} {bn_person_verb(verb, 'perfect', pr)}।",
            "synthetic_new_glossed.eaf",
        )

    def t_coordination():
        pr = pick(PRONOUNS)
        verb = pick([v for v in CORE_VERBS if v["ipa"] in {"ɟɔm", "ɲu", "ɲɛl", "kiɾiɲ"}])
        o1 = pick_object_for(verb)
        o2 = pick_object_for(verb)
        for _ in range(10):
            if o2["ipa"] != o1["ipa"]:
                break
            o2 = pick_object_for(verb)
        return row(
            "coordination",
            [pr["ipa"], "d̪ɔ", o1["ipa"], "aɾ", o2["ipa"], f"{verb['ipa']}-a"],
            [pr["gl"], "TOP", o1["gloss"], "and", o2["gloss"], f"{verb['gloss']}-FIN"],
            f"{subject_start(pr)} will {verb['en']} the {en(o1)} and the {en(o2)}.",
            f"{pr['bn']} {bn(o1)} ও {bn(o2)} {bn_person_verb(verb, 'future', pr)}।",
            "synthetic_bangla_glossed.eaf",
        )

    def t_contrast():
        pr = pick(PRONOUNS)
        verb = pick([v for v in CORE_VERBS if v["type"] == "trans"])
        o1 = pick_object_for(verb)
        o2 = pick_object_for(verb)
        return row(
            "contrast",
            [pr["ipa"], "d̪ɔ", o1["ipa"], f"{verb['ipa']}-a", "mɛnkʰan", o2["ipa"], pr["neg"], f"{verb['ipa']}-a"],
            [pr["gl"], "TOP", o1["gloss"], f"{verb['gloss']}-FIN", "but", o2["gloss"], f"NEG-{pr['gl']}", f"{verb['gloss']}-FIN"],
            f"{subject_start(pr)} will {verb['en']} the {en(o1)}, but {do_aux(pr)} not {verb['en']} the {en(o2)}.",
            f"{pr['bn']} {bn(o1)} {bn_person_verb(verb, 'future', pr)}, কিন্তু {bn(o2)} {bn_person_verb(verb, 'future', pr)} না।",
            "synthetic_dialogues_glossed.eaf",
        )

    def t_imperative():
        verb = pick([v for v in CORE_VERBS if v["ipa"] in {"hiɟuʔ", "calaʔ", "d̪uɽup̚ʔ", "t̪ɘŋgi", "ɟɔm", "ɲu", "paɽha͡o̯", "ɔlɔʔ"}])
        if verb["type"] == "trans":
            obj = pick_object_for(verb)
            return row(
                "imperative",
                ["d̪a.ja", "kat̪ɛ", obj["ipa"], f"{verb['ipa']}-mɛ"],
                ["please", "Manner", obj["gloss"], f"{verb['gloss']}-IMP.2SG.POL"],
                f"Please {verb['en']} the {en(obj)}.",
                f"দয়া করে {bn(obj)} {bn_verb(verb, 'imp')}।",
                "synthetic_imperatives_glossed.eaf",
            )
        place = pick(places_core)
        return row(
            "imperative",
            ["d̪a.ja", "kat̪ɛ", f"{place['ipa']}-ɾɛ", f"{verb['ipa']}-mɛ"],
            ["please", "Manner", f"{place['gloss']}-LOC", f"{verb['gloss']}-IMP.2SG.POL"],
            f"Please {verb['en']} {place_phrase(verb, place)}.",
            f"দয়া করে {bn_loc(place)} {bn_verb(verb, 'imp')}।",
            "synthetic_imperatives_glossed.eaf",
        )

    def t_question():
        pr = pick([p for p in PRONOUNS if p["gl"] in {"1SG", "2SG", "3SG.ANIM", "1PL.EXCL"}])
        obj = pick(CORE_NOUNS)
        variants = [
            (["nit̪ʔ", f"cɛ-ʔ-{pr['agr']}", "cek-a", "ja"], ["now", f"what-NFIN-{pr['gl']}.SUBJ", "do-FIN", "Q"], f"What will {subject_text(pr, lower=True)} do now?", f"{pr['bn']} এখন কী করবে?"),
            (["cɛt̪ʔ-lɛka", f"men-{pr['agr']}-a"], ["how-like", f"be.NFIN-{pr['gl']}-FIN"], f"How {be_aux(pr)} {subject_text(pr, lower=True)}?", f"{pr['bn']} কেমন আছে?"),
            ([obj["ipa"], "d̪ɔ", "ɔka-ɾɛ", "mena-a"], [obj["gloss"], "TOP", "where-LOC", "exist-FIN"], f"Where is the {en(obj)}?", f"{bn(obj)} কোথায় আছে?"),
        ]
        segments, glosses, english, bangla = pick(variants)
        return row("question", segments, glosses, english, bangla, "synthetic_dialogues_glossed.eaf")

    def t_sequence():
        pr = pick(PRONOUNS)
        v1 = pick([v for v in CORE_VERBS if v["ipa"] in {"ɟɔm", "ɲu", "ɔlɔʔ", "paɽha͡o̯", "kiɾiɲ", "ɲɛl"}])
        v2 = pick([v for v in CORE_VERBS if v["ipa"] in {"calaʔ", "hiɟuʔ", "kɘmi", "gat̪e"}])
        obj = pick_object_for(v1)
        place = pick(places_core)
        time = pick([
            {"ipa": "sɛt̪aʔ", "gloss": "morning", "en": "in the morning", "bn": "সকালে"},
            {"ipa": "t̪eheɲ", "gloss": "today", "en": "today", "bn": "আজ"},
            {"ipa": "hɔla", "gloss": "yesterday", "en": "yesterday", "bn": "গতকাল"},
            {"ipa": "ɔkt̪ɛ", "gloss": "time", "en": "at that time", "bn": "সেই সময়"},
        ])
        return row(
            "sequence",
            [f"{time['ipa']}-ɾɛ", pr["ipa"], obj["ipa"], f"{v1['ipa']}-kat̪ɛ", f"{place['ipa']}-ɾɛ", f"{v2['ipa']}-a"],
            [f"{time['gloss']}-LOC", pr["gl"], obj["gloss"], f"{v1['gloss']}-CVB", f"{place['gloss']}-LOC", f"{v2['gloss']}-FIN"],
            f"{cap(time['en'])}, {subject_text(pr, lower=True)} {v1['en']} the {en(obj)} and then {v2['en']} {place_phrase(v2, place)}.",
            f"{time['bn']} {pr['bn']} {bn(obj)} {bn_verb(v1, 'cvb')} {bn_loc(place)} {bn_person_verb(v2, 'future', pr)}।",
            "synthetic_school_dialogue_glossed.eaf",
        )

    def t_classifier_exist():
        noun = pick(humans + [item for item in CORE_NOUNS if item["class"] in {"animal", "object", "food"}])
        place = pick(places_core)
        clf = "mit̪ʔ-ʈɛn" if noun.get("class") in {"human", "animal"} else "mit̪ʔ"
        clf_gl = "one-CLF.HUM" if noun.get("class") in {"human", "animal"} else "one"
        return row(
            "classifier_exist",
            [clf, noun["ipa"], "d̪ɔ", f"{place['ipa']}-ɾɛ", "mena-a"],
            [clf_gl, noun["gloss"], "TOP", f"{place['gloss']}-LOC", "exist-FIN"],
            f"One {en(noun)} is in the {en(place)}.",
            f"একটি {bn(noun)} {bn_loc(place)} আছে।",
            "synthetic_new_glossed_3.eaf",
        )

    templates = [
        (t_topic_adjective, 0.12),
        (t_possessive_adjective, 0.10),
        (t_locative_exist, 0.11),
        (t_transitive_fin, 0.12),
        (t_negative_transitive, 0.10),
        (t_progressive, 0.08),
        (t_past_transitive, 0.08),
        (t_perfect_transitive, 0.06),
        (t_coordination, 0.07),
        (t_contrast, 0.05),
        (t_imperative, 0.05),
        (t_question, 0.03),
        (t_sequence, 0.08),
        (t_classifier_exist, 0.05),
    ]
    return templates


def generate_rows(layout_a, layout_b, parallel, templates, target):
    funcs = [func for func, _ in templates]
    weights = [weight for _, weight in templates]

    existing = set()
    for value in list(layout_a["ref"]) + list(layout_a["mb"]) + list(layout_b["ipa_original"]) + list(layout_b["ipa_normalized"]) + list(parallel["santali"]):
        existing.add(re.sub(r"\s+", " ", (value or "").strip().rstrip(".")))

    synthetic = []
    seen = set()
    attempts = 0
    while len(synthetic) < target and attempts < target * 80:
        attempts += 1
        func = RNG.choices(funcs, weights=weights, k=1)[0]
        try:
            candidate = func()
        except Exception:
            continue
        key = re.sub(r"\s+", " ", candidate["ipa_normalized"].strip().rstrip("."))
        if key in seen or key in existing:
            continue
        if len(candidate["segments"].split(" | ")) != len(candidate["glosses"].split(" | ")):
            continue
        seen.add(key)
        synthetic.append(candidate)

    if len(synthetic) < target:
        raise RuntimeError(f"Only generated {len(synthetic)} unique rows after {attempts} attempts")
    return synthetic, attempts, existing


def write_outputs(synthetic, attempts, existing, layout_a, layout_b, parallel, parallel_raw, text_files, pdf_texts, pdf_meta, dictionary_rows, dictionary_header, semantic_domains, target):
    OUT.mkdir(exist_ok=True)
    layout_b_rows = []
    layout_a_rows = []
    parallel_rows = []
    start = 0

    for idx, item in enumerate(synthetic, start=1):
        seg_count = len(item["segments_list"])
        duration = max(1000, min(6000, 550 + seg_count * 360 + RNG.randint(-120, 220)))
        end = start + duration
        layout_b_rows.append({
            "file": item["file"],
            "utterance_id": str(idx),
            "start_ms": str(start),
            "end_ms": str(end),
            "bangla": f"[SYN.{idx}] {item['bangla_plain']}",
            "english": item["english"],
            "ipa_original": item["ipa_original"],
            "ipa_normalized": item["ipa_normalized"],
            "segments": item["segments"],
            "glosses": item["glosses"],
            "notes": item["notes"],
        })
        layout_a_rows.append({
            "file": item["file"],
            "utterance_id": str(idx),
            "start_ms": str(start),
            "end_ms": str(end),
            "ref": item["ipa_original"].rstrip("."),
            "mb": item["ipa_normalized"].rstrip("."),
            "gl": item["glosses"].replace(" | ", " "),
            "ft": item["english"],
        })
        parallel_rows.append({
            "bangla": item["bangla_plain"],
            "english": item["english"],
            "santali": item["ipa_original"],
        })
        start = end

    outputs = {
        "layoutB": OUT / f"santali_synthetic_layoutB_{target}.csv",
        "layoutA": OUT / f"santali_synthetic_layoutA_{target}.csv",
        "parallel": OUT / f"santali_synthetic_parallel_{target}.csv",
    }
    file_specs = [
        (outputs["layoutB"], layout_b_rows, ["file", "utterance_id", "start_ms", "end_ms", "bangla", "english", "ipa_original", "ipa_normalized", "segments", "glosses", "notes"]),
        (outputs["layoutA"], layout_a_rows, ["file", "utterance_id", "start_ms", "end_ms", "ref", "mb", "gl", "ft"]),
        (outputs["parallel"], parallel_rows, ["bangla", "english", "santali"]),
    ]
    for path, rows, fields in file_specs:
        with path.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=fields)
            writer.writeheader()
            writer.writerows(rows)

    pattern_counts = Counter(item["pattern"] for item in synthetic)
    file_counts = Counter(item["file"] for item in synthetic)
    seg_lengths = [len(item["segments_list"]) for item in synthetic]
    domain_counts = Counter(row.get("Semantic Domain") or "" for row in dictionary_rows)
    pos_counts = Counter(row.get("POS") or "" for row in dictionary_rows)
    gloss_labels = split_glosses(layout_a["gl"]) + split_glosses(layout_b["glosses"])
    pronoun_tokens = re.findall(r"\[[^\]]+\]", pdf_texts.get("Santali_Pronoun.pdf", ""))

    profile = {
        "requested_target": target,
        "ground_truth_only": True,
        "ignored_directory": "generated",
        "files_read_from_ground_truth": sorted(path.name for path in ROOT.iterdir() if path.is_file()),
        "csv_profiles": {
            "output_layoutA (1).csv": {"rows": len(layout_a), "columns": list(layout_a.columns)},
            "output_layoutB (1).csv": {"rows": len(layout_b), "columns": list(layout_b.columns)},
            "Santali_dataset_parsed_bom.csv": {"rows": len(parallel), "columns": list(parallel.columns)},
            "Santali_dataset_parsed.csv": {"rows": len(parallel_raw), "columns": list(parallel_raw.columns)},
        },
        "text_profiles": {
            name: {"chars": len(text), "nonempty_lines": sum(bool(line.strip()) for line in text.splitlines())}
            for name, text in text_files.items()
        },
        "pdf_profiles": {
            name: {"chars_extracted": len(text), **pdf_meta[name]}
            for name, text in pdf_texts.items()
        },
        "xlsx_profiles": {
            "Dictionary": {
                "rows": len(dictionary_rows),
                "columns": dictionary_header,
                "pos_top": pos_counts.most_common(20),
                "semantic_domain_top": domain_counts.most_common(20),
            },
            "Semantic_Domains": {"rows": len(semantic_domains), "sample": semantic_domains[:30]},
        },
        "annotation_label_top": gloss_labels.most_common(60),
        "pronoun_pdf_bracketed_tokens": pronoun_tokens[:80],
        "synthetic_outputs": {key: str(path) for key, path in outputs.items()},
        "synthetic_validation": {
            "rows": len(synthetic),
            "unique_ipa_normalized": len({item["ipa_normalized"] for item in synthetic}),
            "exact_ground_truth_duplicates": sum(1 for item in synthetic if re.sub(r"\s+", " ", item["ipa_normalized"].strip().rstrip(".")) in existing),
            "segments_match_glosses": all(len(item["segments"].split(" | ")) == len(item["glosses"].split(" | ")) for item in synthetic),
            "segment_length_min": min(seg_lengths),
            "segment_length_avg": round(sum(seg_lengths) / len(seg_lengths), 2),
            "segment_length_max": max(seg_lengths),
            "pattern_counts": pattern_counts.most_common(),
            "file_counts": file_counts.most_common(),
            "attempts": attempts,
        },
    }

    summary_path = OUT / f"ground_truth_analysis_summary_{target}.json"
    summary_path.write_text(json.dumps(profile, ensure_ascii=False, indent=2), encoding="utf-8")
    latest_summary_path = OUT / "ground_truth_analysis_summary.json"
    latest_summary_path.write_text(json.dumps(profile, ensure_ascii=False, indent=2), encoding="utf-8")

    readme = [
        "# Santali Synthetic Data - Ground Truth Only",
        "",
        f"Generated {len(synthetic)} synthetic examples from the `ground truth` folder only. The existing `generated` folder was not read or used.",
        "",
        "## Outputs",
        "",
        f"- `{outputs['layoutB'].name}`: rich interlinear-style layout with Bangla, English, IPA, normalized segmentation, glosses, and notes.",
        f"- `{outputs['layoutA'].name}`: compact layout matching `file,utterance_id,start_ms,end_ms,ref,mb,gl,ft`.",
        f"- `{outputs['parallel'].name}`: parallel `bangla,english,santali` layout.",
        f"- `{summary_path.name}`: file profiles, label/domain distributions, extracted pronoun tokens, and validation results.",
        "- `ground_truth_analysis_summary.json`: latest-run copy of the validation summary.",
        "",
        "## Validation",
        "",
        f"- Rows generated: {len(synthetic)}",
        f"- Unique normalized IPA rows: {len({item['ipa_normalized'] for item in synthetic})}",
        f"- Exact duplicates against ground truth Santali/ref/mb strings: {profile['synthetic_validation']['exact_ground_truth_duplicates']}",
        f"- Segment/gloss count alignment: {profile['synthetic_validation']['segments_match_glosses']}",
        f"- Segment length min/avg/max: {min(seg_lengths)} / {round(sum(seg_lengths) / len(seg_lengths), 2)} / {max(seg_lengths)}",
        "",
        "## Pattern Distribution",
        "",
    ]
    readme.extend(f"- {pattern}: {count}" for pattern, count in pattern_counts.most_common())
    readme.extend([
        "",
        "The templates preserve labels and structures observed in the ground truth, including TOP, LOC, GEN, NEG, FIN, PROG, PST, PRF, CVB, IMP, classifier, coordination, contrast, and question patterns.",
        "",
    ])
    (OUT / "README.md").write_text("\n".join(readme), encoding="utf-8")

    return outputs, summary_path, OUT / "README.md", profile


def parse_args():
    parser = argparse.ArgumentParser(description="Generate synthetic Santali examples from ground-truth data only.")
    parser.add_argument(
        "--target",
        type=int,
        default=DEFAULT_TARGET,
        help=f"Number of synthetic examples to generate. Default: {DEFAULT_TARGET}.",
    )
    parser.add_argument(
        "--seed",
        type=int,
        default=20260527,
        help="Random seed for deterministic generation.",
    )
    return parser.parse_args()


def main():
    args = parse_args()
    if args.target < 1:
        raise ValueError("--target must be at least 1")
    global RNG
    RNG = random.Random(args.seed)
    data = read_ground_truth()
    layout_a, layout_b, parallel, parallel_raw, text_files, pdf_texts, pdf_meta, dictionary_rows, dictionary_header, semantic_domains = data
    all_nouns, all_adjs, all_verbs, dict_nouns, _, dict_verbs = build_lexicon(dictionary_rows)
    templates = make_generator(all_nouns, all_adjs, all_verbs, dict_nouns, dict_verbs)
    synthetic, attempts, existing = generate_rows(layout_a, layout_b, parallel, templates, args.target)
    outputs, summary_path, readme_path, profile = write_outputs(
        synthetic,
        attempts,
        existing,
        layout_a,
        layout_b,
        parallel,
        parallel_raw,
        text_files,
        pdf_texts,
        pdf_meta,
        dictionary_rows,
        dictionary_header,
        semantic_domains,
        args.target,
    )
    print(json.dumps({
        "rows": profile["synthetic_validation"]["rows"],
        "outputs": {key: str(value) for key, value in outputs.items()},
        "summary": str(summary_path),
        "readme": str(readme_path),
        "duplicates_against_ground_truth": profile["synthetic_validation"]["exact_ground_truth_duplicates"],
        "segments_match_glosses": profile["synthetic_validation"]["segments_match_glosses"],
        "pattern_counts": profile["synthetic_validation"]["pattern_counts"],
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
